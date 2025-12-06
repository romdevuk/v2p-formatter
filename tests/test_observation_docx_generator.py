"""
Unit tests for Observation DOCX Generator
"""
import unittest
import tempfile
import shutil
from pathlib import Path
from app.observation_docx_generator import create_observation_docx, _extract_placeholders_from_line, _add_media_table_to_doc
from docx import Document


class TestObservationDocxGenerator(unittest.TestCase):
    """Test observation DOCX generator functionality"""
    
    def setUp(self):
        """Set up test fixtures"""
        self.temp_dir = Path(tempfile.mkdtemp())
        self.output_path = self.temp_dir / "test_output.docx"
        
        # Create test image file
        self.test_image = self.temp_dir / "test.jpg"
        self.test_image.write_bytes(b"fake jpg")
    
    def tearDown(self):
        """Clean up test fixtures"""
        shutil.rmtree(self.temp_dir)
    
    def test_extract_placeholders_from_line(self):
        """Test extracting placeholders from a line"""
        from app.observation_docx_generator import _extract_placeholders_from_line
        
        line = "This is {{Placeholder1}} and {{Placeholder2}}"
        placeholders = _extract_placeholders_from_line(line)
        
        self.assertEqual(len(placeholders), 2)
        self.assertIn('placeholder1', placeholders)
        self.assertIn('placeholder2', placeholders)
    
    def test_create_observation_docx_empty(self):
        """Test creating DOCX with empty content"""
        result = create_observation_docx("", {}, self.output_path)
        
        self.assertTrue(result['success'])
        self.assertTrue(self.output_path.exists())
    
    def test_create_observation_docx_with_text(self):
        """Test creating DOCX with text content"""
        text = "This is a test document.\nIt has multiple lines."
        result = create_observation_docx(text, {}, self.output_path)
        
        self.assertTrue(result['success'])
        self.assertTrue(self.output_path.exists())
    
    def test_create_observation_docx_with_placeholders(self):
        """Test creating DOCX with placeholders"""
        text = "This is {{Placeholder1}} and {{Placeholder2}}"
        assignments = {
            'placeholder1': [],
            'placeholder2': []
        }
        
        result = create_observation_docx(text, assignments, self.output_path)
        
        self.assertTrue(result['success'])
        self.assertTrue(self.output_path.exists())
    
    def test_create_observation_docx_with_media(self):
        """Test creating DOCX with assigned media"""
        text = "This is {{Placeholder1}}"
        assignments = {
            'placeholder1': [
                {
                    'path': str(self.test_image),
                    'name': 'test.jpg',
                    'type': 'image'
                }
            ]
        }
        
        result = create_observation_docx(text, assignments, self.output_path)
        
        self.assertTrue(result['success'])
        self.assertTrue(self.output_path.exists())
    
    def test_create_observation_docx_with_video(self):
        """Test creating DOCX with video (filename only)"""
        text = "This is {{Placeholder1}}"
        assignments = {
            'placeholder1': [
                {
                    'path': '/path/to/video.mp4',
                    'name': 'video.mp4',
                    'type': 'video'
                }
            ]
        }
        
        result = create_observation_docx(text, assignments, self.output_path)
        
        self.assertTrue(result['success'])
        self.assertTrue(self.output_path.exists())


if __name__ == '__main__':
    unittest.main()

